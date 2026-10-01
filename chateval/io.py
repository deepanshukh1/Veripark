"""Reading test cases and writing results."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .models import CaseResult, TestCase
from .prompts import PROMPT_VERSION


def _read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def _read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("Reading .xlsx needs openpyxl: pip install openpyxl") from exc
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(h or "").strip() for h in rows[0]]
    return [dict(zip(header, ["" if v is None else str(v) for v in row])) for row in rows[1:]]


def load_cases(path: str | Path, high_risk_keywords: tuple[str, ...] = ()) -> list[TestCase]:
    """CSV/XLSX with columns: question, groundtruth (required); id, source, topic, risk (optional)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = _read_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        raw = _read_xlsx(path)
    else:
        raise ValueError(f"unsupported input type {suffix!r}; use .csv or .xlsx")

    cases: list[TestCase] = []
    for n, row in enumerate(raw, start=1):
        row = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        if not row.get("question"):
            continue
        if not row.get("groundtruth"):
            raise ValueError(f"row {n}: 'groundtruth' is empty (columns found: {sorted(row)})")
        q = row["question"]
        flagged = row.get("risk", "").lower() in ("high", "true", "yes", "1") or any(
            k in q.lower() for k in high_risk_keywords)
        cases.append(TestCase(id=row.get("id") or f"q{n:03d}", question=q, groundtruth=row["groundtruth"],
                              source=row.get("source", ""), topic=row.get("topic", ""), high_risk=flagged))
    if not cases:
        raise ValueError("no test cases found; need columns 'question' and 'groundtruth'")
    return cases


def safe_cell(value) -> str:
    """Neutralise CSV/Excel formula injection from chatbot output (cells starting = + - @)."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


CSV_COLUMNS = ["id", "question", "groundtruth", "bot_answer", "latency_s", "similarity", "judge_score",
               "judge_spread", "judge_issue", "judge_reasoning", "flagged_claims", "verdict", "reason",
               "detail", "high_risk", "prompt_version", "judge_model"]


def _row(r: CaseResult, model: str) -> dict:
    ev = r.evaluation
    return {
        "id": r.case.id, "question": r.case.question, "groundtruth": r.case.groundtruth,
        "bot_answer": r.reply.answer, "latency_s": f"{r.reply.latency_s:.3f}",
        "similarity": "" if ev.similarity is None else f"{ev.similarity:.3f}",
        "judge_score": "" if ev.judge is None else f"{ev.judge.score:.2f}",
        "judge_spread": "" if ev.judge is None else f"{ev.judge.spread:.2f}",
        "judge_issue": "" if ev.judge is None else ev.judge.issue,
        "judge_reasoning": "" if ev.judge is None else ev.judge.reasoning,
        "flagged_claims": "" if ev.hallucination is None else " | ".join(
            f"[{c.status}] {c.text}" for c in ev.hallucination.flagged),
        "verdict": r.decision.verdict.value, "reason": r.decision.reason, "detail": r.decision.detail,
        "high_risk": "yes" if r.case.high_risk else "", "prompt_version": PROMPT_VERSION, "judge_model": model,
    }


def write_csv(results: list[CaseResult], path: Path, model: str) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_COLUMNS)
        writer.writeheader()
        for r in results:
            writer.writerow({k: safe_cell(v) for k, v in _row(r, model).items()})


def write_audit(results: list[CaseResult], path: Path, meta: dict) -> None:
    """One JSON line per judgement with everything needed to reconstruct the decision."""
    with path.open("w", encoding="utf-8") as fh:
        for r in results:
            ev = r.evaluation
            fh.write(json.dumps({
                "run_at": meta["run_at"], "prompt_version": PROMPT_VERSION, "model": meta["chat_model"],
                "thresholds": meta["thresholds"], "case_id": r.case.id, "question": r.case.question,
                "groundtruth": r.case.groundtruth, "bot_answer": r.reply.answer,
                "similarity": ev.similarity,
                "judge": None if ev.judge is None else vars(ev.judge),
                "claims": None if ev.hallucination is None else [vars(c) for c in ev.hallucination.claims],
                "errors": ev.errors, "verdict": r.decision.verdict.value, "reason": r.decision.reason,
            }, ensure_ascii=False) + "\n")
